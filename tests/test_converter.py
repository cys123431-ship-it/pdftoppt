import os
import tempfile
import threading
import unittest

import fitz
from pptx import Presentation

from converter import (
    CANCELLED_MESSAGE,
    CONFLICT_OVERWRITE,
    convert_pdf_to_docx,
    convert_pdf_to_images,
    convert_pdf_to_pptx,
    localize_message,
    merge_pdfs,
    parse_page_range,
)


def make_pdf(path: str, page_sizes=((600, 300),), text="hello") -> None:
    doc = fitz.open()
    try:
        for index, (width, height) in enumerate(page_sizes):
            page = doc.new_page(width=width, height=height)
            page.insert_text((36, 72), f"{text} {index + 1}")
        doc.save(path)
    finally:
        doc.close()


class ConverterTests(unittest.TestCase):
    def test_parse_page_range(self):
        self.assertEqual(parse_page_range("1-3,5,3", 6), [0, 1, 2, 4])
        self.assertEqual(parse_page_range("", 3), [0, 1, 2])
        with self.assertRaises(ValueError):
            parse_page_range("3-2", 4)
        with self.assertRaises(ValueError):
            parse_page_range("0", 4)
        with self.assertRaises(ValueError):
            parse_page_range("5", 4)

    def test_overwrite_failure_preserves_existing_pptx(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = os.path.join(tmp, "source.pdf")
            target = os.path.join(tmp, "existing.pptx")
            make_pdf(source)
            original = b"keep-this-existing-file"
            with open(target, "wb") as file:
                file.write(original)

            success, _message = convert_pdf_to_pptx(
                source,
                target,
                page_range_text="99",
                output_conflict_policy=CONFLICT_OVERWRITE,
            )
            self.assertFalse(success)
            with open(target, "rb") as file:
                self.assertEqual(file.read(), original)

    def test_merge_refuses_to_overwrite_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            first = os.path.join(tmp, "a.pdf")
            second = os.path.join(tmp, "b.pdf")
            make_pdf(first, text="first")
            make_pdf(second, text="second")
            with open(first, "rb") as file:
                original = file.read()

            success, message = merge_pdfs(
                [first, second],
                first,
                output_conflict_policy=CONFLICT_OVERWRITE,
            )
            self.assertFalse(success)
            self.assertIn("cannot overwrite", message)
            with open(first, "rb") as file:
                self.assertEqual(file.read(), original)

    def test_pptx_preserves_mixed_page_aspect_ratio(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = os.path.join(tmp, "mixed.pdf")
            target = os.path.join(tmp, "mixed.pptx")
            make_pdf(source, page_sizes=((600, 300), (300, 600)))

            success, message = convert_pdf_to_pptx(source, target, render_dpi=72)
            self.assertTrue(success, message)
            prs = Presentation(target)
            self.assertEqual(len(prs.slides), 2)
            portrait_picture = prs.slides[1].shapes[0]
            picture_ratio = portrait_picture.width / portrait_picture.height
            self.assertAlmostEqual(picture_ratio, 0.5, delta=0.03)
            self.assertLess(portrait_picture.width, prs.slide_width)
            self.assertEqual(portrait_picture.height, prs.slide_height)

    def test_image_cancellation_rolls_back_staged_outputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = os.path.join(tmp, "source.pdf")
            output_dir = os.path.join(tmp, "images")
            make_pdf(source, page_sizes=((300, 300), (300, 300), (300, 300)))
            event = threading.Event()

            def progress(value: int):
                if value > 0:
                    event.set()

            success, message = convert_pdf_to_images(
                source,
                output_dir,
                image_format="png",
                dpi=72,
                progress_callback=progress,
                cancel_event=event,
            )
            self.assertFalse(success)
            self.assertEqual(message, CANCELLED_MESSAGE)
            self.assertEqual(os.listdir(output_dir), [])

    def test_cancelled_overwrite_keeps_existing_image(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = os.path.join(tmp, "source.pdf")
            output_dir = os.path.join(tmp, "images")
            os.makedirs(output_dir)
            make_pdf(source, page_sizes=((300, 300), (300, 300)))
            existing = os.path.join(output_dir, "source_p001.png")
            original = b"old-image"
            with open(existing, "wb") as file:
                file.write(original)
            event = threading.Event()

            def progress(value: int):
                if value > 0:
                    event.set()

            success, message = convert_pdf_to_images(
                source,
                output_dir,
                image_format="png",
                dpi=72,
                progress_callback=progress,
                output_conflict_policy=CONFLICT_OVERWRITE,
                cancel_event=event,
            )
            self.assertFalse(success)
            self.assertEqual(message, CANCELLED_MESSAGE)
            with open(existing, "rb") as file:
                self.assertEqual(file.read(), original)
            self.assertEqual(sorted(os.listdir(output_dir)), ["source_p001.png"])

    def test_docx_conversion_worker_smoke(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = os.path.join(tmp, "source.pdf")
            target = os.path.join(tmp, "source.docx")
            make_pdf(source)
            success, message = convert_pdf_to_docx(source, target)
            self.assertTrue(success, message)
            self.assertTrue(os.path.isfile(target))
            self.assertGreater(os.path.getsize(target), 0)

    def test_korean_result_localization(self):
        self.assertEqual(
            localize_message(CANCELLED_MESSAGE, "ko"),
            "사용자가 작업을 취소했습니다.",
        )
        self.assertIn(
            "병합",
            localize_message("Merge successful! Saved as: C:/tmp/out.pdf", "ko"),
        )


if __name__ == "__main__":
    unittest.main()
